from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Font

PROJECT = Path(__file__).resolve().parent.parent
RESULTS_DIR = PROJECT / "Results"
OUTPUTS = PROJECT / "outputs"

DST = RESULTS_DIR / "Markify Experiment 2 - Timing & Capacity Planning.xlsx"
RESULTS_CSV = OUTPUTS / "exp2_full_results.csv"
EXTRAP_CSV = OUTPUTS / "exp2_extrapolation.csv"
OLS_CSV = OUTPUTS / "exp2_ols_coefficients.csv"
LATENCY_BY_DB_CSV = OUTPUTS / "exp2_latency_by_db_count.csv"
LATENCY_PNG = OUTPUTS / "exp2_latency_distribution.png"
CONV_PNG = OUTPUTS / "exp2_convergence.png"
PER_DB_PNG = OUTPUTS / "exp2_per_db_latency.png"

SINGLE_LIST_COLS = {
    "scenario_id": "Scenario ID",
    "name": "Candidate Name",
    "db_codes": "Databases Queried",
    "db_count": "# DBs",
    "class_count": "# Classes",
    "total_elapsed": "Total Time (s)",
    "attempt_status": "Outcome",
    "was_split": "Required Split",
}

EXTRAP_COLS = {
    "n_names": "Number of Names",
    "api_calls_needed": "Estimated API Calls",
    "pure_api_hr": "Pure API Time (hours)",
    "quota_periods_needed": "Quota Periods Needed (x 2,000)",
    "ci95_half_sec": "CI95 Half-Width (seconds)",
}

OLS_COLS = {
    "factor": "Factor",
    "coefficient_s": "Coefficient (s per unit)",
    "p_value": "P-value",
    "significant": "Significant?",
}

LATENCY_BY_DB_COLS = {
    "db_count": "# DBs Queried",
    "n_scenarios": "Scenarios",
    "mean_seconds": "Mean (s)",
    "median_seconds": "Median (s)",
    "p95_seconds": "95th %ile (s)",
    "max_seconds": "Max (s)",
}

TITLE = Font(bold=True, size=14)
BOLD = Font(bold=True, size=11)


def write_lines(ws, lines, start_row=1):
    for i, (text, kind) in enumerate(lines):
        cell = ws.cell(start_row + i, 1, text)
        if kind == "title":
            cell.font = TITLE
        elif kind == "bold":
            cell.font = BOLD
    return start_row + len(lines)


def write_table(ws, df, start_row):
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(start_row, col_idx, col_name)
        cell.font = BOLD
    for r_off, row in enumerate(df.itertuples(index=False), 1):
        for c_off, val in enumerate(row, 1):
            ws.cell(start_row + r_off, c_off, val)
    return start_row + len(df)


def add_sized_image(ws, path, anchor_row, width=720, height=340):
    img = Image(str(path))
    img.width = width
    img.height = height
    ws.add_image(img, ws.cell(anchor_row, 1).coordinate)
    return anchor_row + (height // 20) + 2


def compute_metadata():
    df = pd.read_csv(RESULTS_CSV)
    n = len(df)
    valid_mask = df["attempt_status"].isin(["ok", "split_ok", "split_partial"])
    valid_df = df[valid_mask]
    n_valid = len(valid_df)
    n_ok = int((df["attempt_status"] == "ok").sum())
    n_split_ok = int((df["attempt_status"] == "split_ok").sum())
    n_split_partial = int((df["attempt_status"] == "split_partial").sum())
    n_failed = n - n_valid
    avg_time = float(valid_df["total_elapsed"].mean()) if n_valid else 0.0
    return {
        "total": n,
        "valid": n_valid,
        "ok": n_ok,
        "split_ok": n_split_ok,
        "split_partial": n_split_partial,
        "failed": n_failed,
        "avg_time": avg_time,
        "coverage_pct": n / 756 * 100,
    }


def write_summary(wb):
    ws = wb.create_sheet("Summary")
    meta = compute_metadata()

    top_factor_text = "(OLS results not available)"
    if OLS_CSV.exists():
        ols = pd.read_csv(OLS_CSV)
        sig = ols[ols["significant"] == "Yes"].sort_values("coefficient_s", ascending=False)
        if len(sig) > 0:
            top = sig.iloc[0]
            top_factor_text = f"{top['factor']}  (+{top['coefficient_s']:.1f}s)"

    quota_days_20k_text = "(extrapolation not available)"
    if EXTRAP_CSV.exists():
        ex = pd.read_csv(EXTRAP_CSV)
        row20k = ex[ex["n_names"] == 20000]
        if len(row20k) > 0:
            quota_days_20k_text = f"{row20k.iloc[0]['quota_periods_needed']:.0f} days at 2,000 calls/day"

    guide = [
        ("EXPERIMENT 2 - TIMING & CAPACITY PLANNING", "title"),
        ("", None),
        ("PURPOSE", "bold"),
        ("Project Markify batch runtime (500-20k names) and identify what slows it down.", None),
        ("", None),
        ("EXPERIMENT METADATA", "bold"),
        (f"  Scenarios collected:     {meta['total']} of 756 ({meta['coverage_pct']:.1f}%)", None),
        (f"  Valid timing data:       {meta['valid']}", None),
        (f"  Single-call success:     {meta['ok']}", None),
        (f"  Split recovery (ok):     {meta['split_ok']}", None),
        (f"  Split partial recovery:  {meta['split_partial']}", None),
        (f"  Failed (timeout / etc):  {meta['failed']}", None),
        (f"  Average time/scenario:   {meta['avg_time']:.2f}s", None),
        ("", None),
        ("DESIGN", "bold"),
        ("  Main run:     756 scenarios timed once each   -> mu (mean)", None),
        ("  Variance run: 10 scenarios timed 30x          -> sigma (stability)", None),
        ("  Combined: batch time = mu * N +/- margin", None),
        ("", None),
        ("WHAT EACH TAB SHOWS", "bold"),
        ("  Tab 1   Summary           this tab - design + key findings", None),
        ("  Tab 2   Average Timing    histogram + breakdown by # DBs", None),
        ("  Tab 3   Stability         convergence chart for repeated measurements", None),
        ("  Tab 4   Extrapolation     projection table for 500-20k names", None),
        ("  Tab 5   Factors           which DB / factor slows calls down", None),
        ("  Tab 6   Raw Data          every scenario's individual timing + status", None),
        ("", None),
        ("HOW TO READ", "bold"),
        ("  Tab 2: histogram shape -> where typical calls land", None),
        ("  Tab 3: flat line = stable measurement; drifting = need more rounds", None),
        ("  Tab 4: read the row for your batch size", None),
        ("  Tab 5: bigger coefficient = bigger latency impact", None),
        ("  Tab 6: filter Outcome to see only successful runs", None),
        ("", None),
        ("KEY FINDINGS", "bold"),
        (f"  Average call: {meta['avg_time']:.1f}s (high variance with parameters)", None),
        (f"  Top latency driver: {top_factor_text}", None),
        (f"  20k names projected: ~{quota_days_20k_text}", None),
        ("  REAL bottleneck is quota, not latency.", None),
    ]
    write_lines(ws, guide)
    ws.column_dimensions["A"].width = 100


def write_average_timing(wb):
    ws = wb.create_sheet("Average Timing")

    intro = [
        ("AVERAGE TIMING - how long does a typical call take?", "title"),
        ("", None),
        ("HOW TO READ", "bold"),
        ("  Histogram below: X = time per scenario (seconds), Y = number of scenarios", None),
        ("  Two colors: blue = single call worked; orange = had to split per DB", None),
        ("  Table below shows breakdown by # of databases queried.", None),
        ("", None),
    ]
    next_row = write_lines(ws, intro)

    if LATENCY_PNG.exists():
        next_row = add_sized_image(ws, LATENCY_PNG, next_row)

    table_intro = [
        ("BREAKDOWN BY # OF DATABASES", "bold"),
        ("  More DBs = longer calls. This table shows the trend numerically.", None),
        ("", None),
    ]
    next_row = write_lines(ws, table_intro, start_row=next_row + 2)

    if LATENCY_BY_DB_CSV.exists():
        df = pd.read_csv(LATENCY_BY_DB_CSV)
        keep = [c for c in LATENCY_BY_DB_COLS if c in df.columns]
        df = df[keep].rename(columns={c: LATENCY_BY_DB_COLS[c] for c in keep})
        write_table(ws, df, start_row=next_row)

    widths = [16, 14, 14, 14, 16, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


def write_stability(wb):
    ws = wb.create_sheet("Stability")

    intro = [
        ("STABILITY - are repeated measurements consistent?", "title"),
        ("", None),
        ("HOW TO READ", "bold"),
        ("  Each line = one scenario, repeated up to 30 times.", None),
        ("  X = round number; Y = running mean of time so far.", None),
        ("  Flat line = sigma is stable (good).", None),
        ("  Drifting line = need more rounds to stabilize.", None),
        ("", None),
        ("WHY IT MATTERS", "bold"),
        ("  We extrapolate batch runtimes using mu (mean) plus a sigma-based margin.", None),
        ("  Without a stable sigma, the +/- margin in Tab 4 cannot be trusted.", None),
        ("", None),
    ]
    next_row = write_lines(ws, intro)
    if CONV_PNG.exists():
        add_sized_image(ws, CONV_PNG, next_row)

    ws.column_dimensions["A"].width = 100


def write_extrapolation(wb):
    ws = wb.create_sheet("Extrapolation")

    intro = [
        ("EXTRAPOLATION - projected runtime for batch jobs", "title"),
        ("", None),
        ("HOW TO READ", "bold"),
        ("  Estimated API Calls = N x avg calls per name (including split sub-calls).", None),
        ("  Pure API Time = wall-clock if quota were unlimited.", None),
        ("  Quota Periods Needed = total calls / 2,000 (Markify daily cap).", None),
        ("  CI95 Half-Width = +/- range from variance subset (1.96 * sigma * sqrt(N)).", None),
        ("", None),
        ("KEY FINDING", "bold"),
        ("  Quota dominates: 20k names ~ ~30 days, NOT 24 hours of API time.", None),
        ("  Higher Markify tier required for production-scale batches.", None),
        ("", None),
    ]
    next_row = write_lines(ws, intro)

    if EXTRAP_CSV.exists():
        df = pd.read_csv(EXTRAP_CSV)
        keep = [c for c in EXTRAP_COLS if c in df.columns]
        df = df[keep].rename(columns={c: EXTRAP_COLS[c] for c in keep})
        write_table(ws, df, start_row=next_row)

    widths = [22, 25, 22, 32, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


def write_factors(wb):
    ws = wb.create_sheet("Factors")

    intro = [
        ("FACTORS - what makes calls slow?", "title"),
        ("", None),
        ("CHART: Mean time per database (single-DB calls)", "bold"),
        ("  Each bar = one database's mean response time.", None),
        ("  Sorted slowest -> fastest. Slowest highlighted in crimson.", None),
        ("", None),
    ]
    next_row = write_lines(ws, intro)

    if PER_DB_PNG.exists():
        next_row = add_sized_image(ws, PER_DB_PNG, next_row, width=720, height=400)

    table_intro = [
        ("OLS REGRESSION: Total Time ~ # DBs + # Classes + Name Length + Name Type", "bold"),
        ("  Coefficient = predicted seconds added per unit increase in that factor.", None),
        ("  P-value < 0.05 means statistically significant.", None),
        ("", None),
    ]
    next_row = write_lines(ws, table_intro, start_row=next_row + 2)

    if OLS_CSV.exists():
        df = pd.read_csv(OLS_CSV)
        keep = [c for c in OLS_COLS if c in df.columns]
        df = df[keep].rename(columns={c: OLS_COLS[c] for c in keep})
        write_table(ws, df, start_row=next_row)

    widths = [50, 24, 12, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


def write_raw_data(wb):
    ws = wb.create_sheet("Raw Data")

    intro = [
        ("RAW DATA - every scenario's individual timing + status", "title"),
        ("", None),
        ("HOW TO READ", "bold"),
        ("  Each row = one scenario from Sheet 2.", None),
        ("  Outcome: ok = single call worked; split_ok / split_partial = needed per-DB split.", None),
        ("  Filter Outcome for successful runs only, or by # DBs to focus on a complexity tier.", None),
        ("", None),
    ]
    next_row = write_lines(ws, intro)

    df = pd.read_csv(RESULTS_CSV)
    keep = [c for c in SINGLE_LIST_COLS if c in df.columns]
    df = df[keep].rename(columns={c: SINGLE_LIST_COLS[c] for c in keep})
    header_row = next_row
    write_table(ws, df, start_row=header_row)

    widths = [14, 22, 35, 8, 10, 14, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate


def main():
    if not RESULTS_CSV.exists():
        raise FileNotFoundError(f"Missing: {RESULTS_CSV}")

    RESULTS_DIR.mkdir(exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_summary(wb)
    write_average_timing(wb)
    write_stability(wb)
    write_extrapolation(wb)
    write_factors(wb)
    write_raw_data(wb)

    wb.save(DST)
    print(f"Built: {DST}")
    print(f"Tabs: {wb.sheetnames}")


if __name__ == "__main__":
    main()
